"""DGP workbook → construction types and work trees.

xlsx is a reference of possible works per construction type.
It is not a calendar schedule and never becomes ExpectedState by itself.
"""

from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from functools import lru_cache
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from sitewatch.settings import project_root

TYPE_COLUMNS: tuple[tuple[str, str], ...] = (
    ("housing", "Жильё"),
    ("education", "Образование"),
    ("healthcare", "Здравоохранение"),
    ("sport", "Спорт"),
    ("culture", "Культура"),
    ("admin", "Административные здания"),
    ("kindergarten", "ДОУ"),
    ("office", "Офисно-деловой центр"),
    ("roads", "Дороги"),
)

DEFAULT_XLSX = (
    project_root()
    / "ТЗ "
    / "Датасет "
    / "7.ДГП_датасеты"
    / "Сводный перечень строительных работ_ЛТЦ.xlsx"
)

_SECTION_FILL = "D9D9D9"
_ID_SAFE = re.compile(r"[^a-z0-9_]+")
_SECTION_CODE = re.compile(r"^\d+\.?$")


def _xlsx_path() -> Path:
    raw = (os.environ.get("SITEWATCH_CONSTRUCTION_XLSX") or "").strip()
    return Path(raw) if raw else DEFAULT_XLSX


def _cell_fill_rgb(cell) -> str:
    fill = cell.fill
    if not fill or not fill.fgColor:
        return ""
    color = fill.fgColor
    rgb = getattr(color, "rgb", None)
    if not rgb:
        return ""
    text = str(rgb).upper()
    if len(text) == 8 and text.startswith("FF"):
        return text[2:]
    return text


def _is_checked(value: Any) -> bool:
    if value is None:
        return False
    text = str(value).strip()
    return bool(text)


def normalize_work_code(raw: Any) -> str | None:
    """Recover spreadsheet codes; datetime cells are corrupted codes, not dates."""
    if raw is None:
        return None
    if isinstance(raw, datetime):
        return f"{raw.day}.{raw.month:02d}"
    if isinstance(raw, date) and not isinstance(raw, datetime):
        return f"{raw.day}.{raw.month:02d}"
    if isinstance(raw, int):
        return str(raw)
    if isinstance(raw, float):
        if raw.is_integer():
            return str(int(raw))
        return str(raw).strip() or None
    text = str(raw).strip()
    return text or None


def _slug_from_code(code: str) -> str:
    text = code.strip().lower().replace(".", "_")
    text = _ID_SAFE.sub("_", text)
    text = re.sub(r"_+", "_", text).strip("_")
    return text


def _work_id(code: str | None, *, parent_id: str | None, orphan_index: int) -> str:
    if code:
        return f"w_{_slug_from_code(code)}"
    if parent_id:
        return f"{parent_id}_d{orphan_index}"
    return f"w_orphan_d{orphan_index}"


def _section_id(code: str | None, name: str, index: int) -> str:
    if code:
        slug = _slug_from_code(code.rstrip("."))
        if slug:
            return f"sec_{slug}"
    slug = _slug_from_code(name) or f"section_{index}"
    return f"sec_{slug}"


def _unique_id(candidate: str, used: set[str]) -> str:
    if candidate not in used:
        used.add(candidate)
        return candidate
    n = 2
    while f"{candidate}_{n}" in used:
        n += 1
    final = f"{candidate}_{n}"
    used.add(final)
    return final


@dataclass
class _Node:
    kind: str  # section | work | detail
    code: str | None
    name: str
    flags: dict[str, bool]
    id: str = ""
    children: list[_Node] = field(default_factory=list)


@dataclass
class _Catalog:
    types: list[dict[str, str]]
    nodes: list[_Node]


def _is_section(code: str | None, *, fill_rgb: str, bold: bool) -> bool:
    if fill_rgb == _SECTION_FILL:
        return True
    if bold and code and _SECTION_CODE.match(code.strip()):
        return True
    return False


def _parse_workbook(path: Path) -> _Catalog:
    if not path.is_file():
        raise FileNotFoundError(f"construction catalog not found: {path}")
    wb = load_workbook(path, data_only=False, read_only=False)
    try:
        ws = wb.active
        header_row = next(ws.iter_rows(min_row=3, max_row=3))
        headers = [cell.value for cell in header_row]
        type_cols: list[tuple[str, str, int]] = []
        for type_id, type_name in TYPE_COLUMNS:
            try:
                col_idx = next(i for i, value in enumerate(headers) if value == type_name)
            except StopIteration as exc:
                raise ValueError(f"construction catalog missing column {type_name!r}") from exc
            type_cols.append((type_id, type_name, col_idx))

        used_ids: set[str] = set()
        roots: list[_Node] = []
        current_section: _Node | None = None
        current_work: _Node | None = None
        orphan_counters: dict[str, int] = {}

        for row in ws.iter_rows(min_row=4, max_row=ws.max_row):
            name_cell = row[1]
            name_raw = name_cell.value
            if name_raw is None or not str(name_raw).strip():
                continue
            name = str(name_raw).strip()
            code = normalize_work_code(row[0].value)
            fill_rgb = _cell_fill_rgb(name_cell)
            bold = bool(name_cell.font and name_cell.font.bold)
            flags = {type_id: _is_checked(row[col_idx].value) for type_id, _, col_idx in type_cols}

            if _is_section(code, fill_rgb=fill_rgb, bold=bold):
                node_id = _unique_id(_section_id(code, name, len(roots) + 1), used_ids)
                section = _Node(kind="section", code=code, name=name, flags=flags, id=node_id)
                roots.append(section)
                current_section = section
                current_work = None
                continue

            if bold:
                parent_id = current_section.id if current_section else None
                node_id = _unique_id(_work_id(code, parent_id=parent_id, orphan_index=1), used_ids)
                work = _Node(kind="work", code=code, name=name, flags=flags, id=node_id)
                if current_section is not None:
                    current_section.children.append(work)
                else:
                    roots.append(work)
                current_work = work
                continue

            parent = current_work or current_section
            parent_key = parent.id if parent else "_root"
            orphan_counters[parent_key] = orphan_counters.get(parent_key, 0) + 1
            parent_id = parent.id if parent else None
            node_id = _unique_id(
                _work_id(code, parent_id=parent_id, orphan_index=orphan_counters[parent_key]),
                used_ids,
            )
            detail = _Node(kind="detail", code=code, name=name, flags=flags, id=node_id)
            if current_work is not None:
                current_work.children.append(detail)
            elif current_section is not None:
                # Detail without a bold parent work — treat as work under the section.
                detail.kind = "work"
                current_section.children.append(detail)
                current_work = detail
            else:
                roots.append(detail)

        return _Catalog(
            types=[{"id": type_id, "name": type_name} for type_id, type_name, _ in type_cols],
            nodes=roots,
        )
    finally:
        wb.close()


@lru_cache(maxsize=1)
def _load_catalog() -> _Catalog:
    return _parse_workbook(_xlsx_path())


def clear_construction_catalog_cache() -> None:
    _load_catalog.cache_clear()


def list_construction_types() -> list[dict[str, str]]:
    return [dict(item) for item in _load_catalog().types]


def known_construction_type_ids() -> set[str]:
    return {item["id"] for item in _load_catalog().types}


def construction_type_name(type_id: str | None) -> str | None:
    if not type_id:
        return None
    for item in _load_catalog().types:
        if item["id"] == type_id:
            return item["name"]
    return None


def _node_applies(node: _Node, type_id: str) -> bool:
    if node.flags.get(type_id):
        return True
    return any(_node_applies(child, type_id) for child in node.children)


def _serialize_work(node: _Node, type_id: str) -> dict[str, Any] | None:
    if not _node_applies(node, type_id):
        return None
    children = []
    for child in node.children:
        if child.kind == "detail":
            if child.flags.get(type_id) or _node_applies(child, type_id):
                children.append(
                    {
                        "id": child.id,
                        "code": child.code,
                        "name": child.name,
                    }
                )
        else:
            nested = _serialize_work(child, type_id)
            if nested is not None:
                # Flatten unexpected nesting into children list shape.
                children.append(
                    {
                        "id": nested["id"],
                        "code": nested["code"],
                        "name": nested["name"],
                    }
                )
    if not node.flags.get(type_id) and not children:
        return None
    return {
        "id": node.id,
        "code": node.code,
        "name": node.name,
        "children": children,
    }


def list_works_for_type(type_id: str) -> dict[str, Any]:
    catalog = _load_catalog()
    type_meta = next((item for item in catalog.types if item["id"] == type_id), None)
    if type_meta is None:
        raise KeyError("construction type not found")
    groups: list[dict[str, Any]] = []
    misc_works: list[dict[str, Any]] = []
    for node in catalog.nodes:
        if node.kind != "section":
            work = _serialize_work(node, type_id)
            if work is not None:
                misc_works.append(work)
            continue
        if not _node_applies(node, type_id):
            continue
        works: list[dict[str, Any]] = []
        for child in node.children:
            serialized = _serialize_work(child, type_id)
            if serialized is not None:
                works.append(serialized)
        if not works:
            continue
        groups.append(
            {
                "id": node.id,
                "code": node.code,
                "name": node.name,
                "works": works,
            }
        )
    if misc_works:
        groups.append(
            {
                "id": "sec_misc",
                "code": None,
                "name": "Прочее",
                "works": misc_works,
            }
        )
    return {"type": dict(type_meta), "groups": groups}
